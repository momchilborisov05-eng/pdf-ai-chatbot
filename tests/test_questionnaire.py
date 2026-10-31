"""
Tests for questionnaire autofill: importing the answer library, answers that
must cite real library entries, reuse until the library changes, the Excel
export, and the folder watcher. No API calls.
"""

import io
import sys
import uuid
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import autofill  # noqa: E402
from questionnaire import LibraryError, ground_answers, parse_library, select_entries  # noqa: E402
from samples.build import text_pdf  # noqa: E402
from storage import Storage  # noqa: E402
from test_api import ALICE, BOB, events, upload  # noqa: E402
from test_review import review_message  # noqa: E402

LIBRARY_CSV = (
    "question;answer;source\n"
    "Are you ISO 27001 certified?;Yes, certified by Bureau Veritas until 2027-03-14.;ISO certificate\n"
    "Where is data stored?;In the EU only (Frankfurt and Amsterdam).;Data residency statement\n"
    ";an answer without a question is skipped;\n"
).encode()

PAGES = ["SUPPLIER QUESTIONNAIRE\n1.1 Is your company certified to ISO/IEC 27001?\n"
         "1.2 In which countries is our data stored?\n1.3 Do you carry cyber insurance?"]


def questionnaire_pdf(folder: Path) -> Path:
    path = folder / f"test-{uuid.uuid4().hex[:8]}.pdf"
    path.write_bytes(text_pdf([PAGES[0] + f"\nRef {uuid.uuid4().hex}"]))
    return path


def sheet(ids: list[int]) -> dict:
    """What the model returns; `ids` are the library entry ids it cites for questions 1.1 and 1.2."""
    return {"title": "Supplier questionnaire", "questions": [
        {"number": "1.1", "question": "Is your company certified to ISO/IEC 27001?", "page": 1,
         "status": "answered", "answer": "Yes, certified by Bureau Veritas until 2027-03-14.",
         "sources": [ids[0]], "note": ""},
        {"number": "1.2", "question": "In which countries is our data stored?", "page": 1,
         "status": "answered", "answer": "Only in the EU.", "sources": [ids[1]], "note": ""},
        {"number": "1.3", "question": "Do you carry cyber insurance?", "page": 1,
         "status": "answered", "answer": "Yes, EUR 5 million.", "sources": [], "note": ""},  # invented
    ]}


def import_csv(http, data=LIBRARY_CSV, headers=None):
    return http.post("/api/library", files={"file": ("answers.csv", data, "text/csv")}, headers=headers or {})


# ---------- the answer library ----------

def test_library_is_read_from_csv_with_any_delimiter_and_bulgarian_headers():
    entries = parse_library("lib.csv", "въпрос,отговор\nИмате ли ISO 27001?,Да.\n".encode("cp1251"))
    assert entries == [{"question": "Имате ли ISO 27001?", "answer": "Да.", "source": ""}]

    entries = parse_library("lib.csv", LIBRARY_CSV)
    assert [e["source"] for e in entries] == ["ISO certificate", "Data residency statement"]


def test_library_is_read_from_excel_without_a_header_row():
    book = Workbook()
    book.active.append(["Do you have a DPO?", "Yes, Ivan Petrov.", "Org chart"])
    out = io.BytesIO()
    book.save(out)

    assert parse_library("lib.xlsx", out.getvalue()) == [
        {"question": "Do you have a DPO?", "answer": "Yes, Ivan Petrov.", "source": "Org chart"}]


@pytest.mark.parametrize("name,data,message", [
    ("lib.pdf", b"%PDF-", r"Excel \(\.xlsx\) or CSV"),
    ("lib.csv", b"question,answer\n", "No rows"),
    ("lib.xlsx", b"not a workbook", "could not be read"),
])
def test_unusable_library_files_are_rejected(name, data, message):
    with pytest.raises(LibraryError, match=message):
        parse_library(name, data)


def test_a_large_library_sends_the_entries_closest_to_the_document():
    entries = [{"id": i, "question": f"Unrelated topic number {i} about catering", "answer": "x" * 500}
               for i in range(50)] + [{"id": 99, "question": "Is the company certified?", "answer": "Yes."}]
    chosen = select_entries(entries, ["Is your company certified to ISO 27001?"], budget=2000)
    assert 99 in [e["id"] for e in chosen]
    assert len(chosen) < len(entries)


# ---------- grounding ----------

def test_answers_without_a_real_library_source_are_not_delivered():
    library = {1: {"question": "ISO?", "source": "ISO certificate"}, 2: {"question": "Data?", "source": ""}}
    report = sheet([1, 77])  # 77 doesn't exist

    counts = ground_answers(report, library)

    iso, data, insurance = report["questions"]
    assert iso["status"] == "answered" and iso["source_refs"][0]["source"] == "ISO certificate"
    assert data["status"] == "needs_input" and data["answer"] == "" and data["draft"] == "Only in the EU."
    assert insurance["status"] == "needs_input" and "No library entry" in insurance["note"]
    assert counts == {"answered": 1, "partial": 0, "needs_input": 2, "unsupported": 2}


# ---------- API ----------

def test_questionnaire_is_answered_from_the_library_and_reused_until_it_changes(make_client, tmp_path):
    http, fake = make_client()
    import_csv(http)
    ids = [e["id"] for e in http.get("/api/library").json()["entries"]]
    fake.queue += [review_message(sheet(ids)), review_message(sheet(ids))]
    doc = upload(http, questionnaire_pdf(tmp_path)).json()
    url = f"/api/documents/{doc['id']}/autofill"

    name, run = events(http.post(url, json={"language": "Bulgarian"}))[-1]

    assert name == "done" and run["playbook"] == "questionnaire"
    report = run["report"]
    assert report["grounding"] == {"answered": 2, "partial": 0, "needs_input": 1, "unsupported": 1}
    assert report["verification"]["verified"] == 3  # every question is in the PDF
    request = fake.requests[0]
    library_text = request["messages"][0]["content"][0]
    assert f'<entry id="{ids[0]}" source="ISO certificate">' in library_text["text"]
    assert library_text["cache_control"] == {"type": "ephemeral"}
    assert "citations" not in request["messages"][0]["content"][1]

    assert http.post(url, json={"language": "Bulgarian"}).json()["cached"] is True
    assert len(fake.requests) == 1

    import_csv(http)  # a new library: the saved answers are out of date
    assert events(http.post(url, json={"language": "Bulgarian"}))[-1][0] == "done"
    assert len(fake.requests) == 2


def test_questionnaire_needs_a_library(make_client, tmp_path):
    http, fake = make_client()
    doc = upload(http, questionnaire_pdf(tmp_path)).json()
    response = http.post(f"/api/documents/{doc['id']}/autofill", json={})
    assert response.status_code == 409 and "library is empty" in response.json()["detail"]
    assert fake.requests == []


def test_answer_sheet_downloads_as_excel_and_csv(make_client, tmp_path):
    http, fake = make_client()
    import_csv(http)
    ids = [e["id"] for e in http.get("/api/library").json()["entries"]]
    report = sheet(ids)
    report["questions"][0]["answer"] = "=HYPERLINK(\"http://evil\")"
    fake.queue.append(review_message(report))
    doc = upload(http, questionnaire_pdf(tmp_path)).json()
    run = events(http.post(f"/api/documents/{doc['id']}/autofill", json={}))[-1][1]

    xlsx = http.get(f"/api/reviews/{run['id']}/answers.xlsx")
    rows = list(load_workbook(io.BytesIO(xlsx.content)).active.iter_rows(values_only=True))
    assert rows[0][:5] == ("No.", "Question", "Page", "Status", "Answer")
    assert rows[1][4] == "=HYPERLINK(\"http://evil\")"  # stored as text, not a formula
    assert rows[1][5].endswith("ISO certificate")
    assert rows[3][3] == "Needs input"

    csv_text = http.get(f"/api/reviews/{run['id']}/export.csv").text
    assert "Needs input" in csv_text and "'=HYPERLINK" in csv_text


def test_libraries_and_answer_sheets_are_private(make_client, tmp_path):
    http, fake = make_client(auth_mode="easyauth")
    import_csv(http, headers=ALICE)
    ids = [e["id"] for e in http.get("/api/library", headers=ALICE).json()["entries"]]
    fake.queue.append(review_message(sheet(ids)))
    doc = upload(http, questionnaire_pdf(tmp_path), headers=ALICE).json()
    run = events(http.post(f"/api/documents/{doc['id']}/autofill", json={}, headers=ALICE))[-1][1]

    assert http.get("/api/library", headers=BOB).json()["count"] == 0
    assert http.get(f"/api/reviews/{run['id']}/answers.xlsx", headers=BOB).status_code == 404


# ---------- the folder watcher ----------

def test_watcher_answers_every_pdf_in_the_folder_and_moves_it(tmp_path, monkeypatch):
    monkeypatch.setenv("DEMO_MODE", "1")
    db = f"sqlite:///{tmp_path / 'watch.db'}"
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    (inbox / "good.pdf").write_bytes(text_pdf(PAGES))
    (inbox / "broken.pdf").write_bytes(b"not a pdf")
    library = tmp_path / "answers.csv"
    library.write_bytes(LIBRARY_CSV)

    autofill.main(["--db", db, "import", str(library)])
    autofill.main(["--db", db, "watch", str(inbox), "--once"])

    assert sorted(p.name for p in (inbox / "done").iterdir()) == ["good.pdf"]
    assert "not a PDF" in (inbox / "failed" / "broken.txt").read_text()
    rows = list(load_workbook(inbox / "answers" / "good-answers.xlsx").active.iter_rows(values_only=True))
    status = {r[0]: r[3] for r in rows[1:]}
    assert status["1.1"] == "Answered" and status["1.3"] == "Needs input"  # insurance isn't in the library
    store = Storage(db)
    assert store.list_reviews(store.list_documents()[0]["id"])[0]["playbook"] == "questionnaire"
    store.close()


def test_sample_questionnaire_is_up_to_date():
    from samples import build

    assert (build.HERE / "supplier-questionnaire.pdf").read_bytes() == text_pdf(build.QUESTIONNAIRE)
